"""Build the mesh video deck from slides/figures/ (made by slides/make_figures.py):
  slides/mesh_ZollerQu.pptx (python-pptx), slides/mesh_ZollerQu.pdf and slides/speaker_notes.md.
One layout description is rendered twice: into the .pptx and, with matplotlib, into the PDF (same positions, sizes, fonts and
explicit line breaks), so the PDF needs no office program. Every number on a slide is read from slides/figures/numbers.json,
which make_figures.py fills from the result files.
usage: pip install python-pptx ; python slides/make_figures.py ; python slides/make_slides.py
"""
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyBboxPatch
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SLIDES, FIG = ROOT / "slides", ROOT / "slides" / "figures"
N = json.loads((FIG / "numbers.json").read_text())["numbers"]
SOURCES = json.loads((FIG / "numbers.json").read_text())["figure_sources"]
W, H = 13.333, 7.5
# calm palette in the direction of the IonQ and Qollab sites (near-black, off-white, neutral greys, one soft orange); no logos.
# ACCENT = soft blue for labels; HIGHLIGHT = soft orange, used ONLY for our IonQ noise-model result.
BG, PANEL, WHITE, MUTED, ACCENT, HIGHLIGHT = "#0F161B", "#1A242B", "#FAFAFA", "#9AA3AA", "#9FB4D1", "#F89B33"
FOOTER = "ZollerQu - Qollab x IonQ Hackathon 2026"
ATTRIBUTION = "This effort is supported by Qollab & IonQ"
LINE = 1.2                                   # line height as a multiple of the font size (both renderers)


def text(x, y, w, content, size=18, colour=WHITE, bold=False, italic=False, align="l", font="Arial"):
    return dict(kind="text", x=x, y=y, w=w, lines=content.split("\n"), size=size, colour=colour, bold=bold, italic=italic, align=align, font=font)


def image(name, x, y, w, h, alt):
    return dict(kind="image", path=FIG / name, x=x, y=y, w=w, h=h, alt=alt)


def card(x, y, w, h):
    return dict(kind="card", x=x, y=y, w=w, h=h)


def source(x, y, w, content):
    return text(x, y, w, content, size=14, colour=MUTED, italic=True)


def pct(value, digits=1):
    return f"{100 * value:.{digits}f} %"


# ---------------------------------------------------------------- numbers used in the slide text (all from the result files) ----
p_ionq, p_ideal, p_random = N["ionq_p_opt"], N["noiseless_p_opt"], N["random_p_opt"]
gain = round(N["gain_over_random"], -1)
loss_points = 100 * (p_ideal - p_ionq)
estimated_loss = 100 * (p_ideal - N["crude_estimate_p_opt"])
date = "11 Oct 2026"
assert N["ionq_date"] == "2026-10-11" and N["ionq_most_frequent_is_optimum"]
g12, g29 = N["gates_12"], N["gates_29"]

ALT = {
    "fig_problem.png": f"An L-shaped triangle mesh split in two parts. The large teal part is computed by the GPU ({pct(N['gpu_share_of_optimum'], 0)} of the work), "
                       f"the small light part by the CPU. An orange line marks the border between the parts.",
    "fig_formula.png": "The Ising energy: a sum over neighbouring regions of border cost times (1 - z_i z_j) / 2, plus lambda times the squared difference "
                       "between the weighted sum of the z_i and (1 - 2 alpha). The first term is the border that is cut, the second a penalty if the work share is off target.",
    "fig_circuit.png": "Circuit sketch with four of the twelve qubit lines: an H gate on every qubit, a cost layer with one two-qubit rzz gate per pair of "
                       "regions, an rx mixer gate on every qubit, then measurement. The cost layer and the mixer are repeated p times.",
    "fig_gates.png": f"Bar chart of two-qubit gate counts. One QAOA layer at 12 qubits: {g12['all_to_all']} on all-to-all connectivity, {g12['heavy_hex']:.0f} routed on "
                     f"heavy-hex. The varQITE ZY circuit at 12 qubits: {N['gates_zy']['all_to_all']} against {N['gates_zy']['heavy_hex']:.0f}. One QAOA layer at 29 qubits: "
                     f"{g29['all_to_all']} against {g29['heavy_hex']:,.0f}.",
    "fig_ionq_result.png": f"Bar chart on a log scale of the share of shots that are the exact best split: random guessing {pct(p_random, 3)}, "
                           f"IonQ Forte 1 noise model {pct(p_ionq)}, noiseless simulation {pct(p_ideal)}. A dashed line marks our estimate of {pct(N['crude_estimate_p_opt'], 0)}.",
    "fig_scaling.png": "Two line charts, for two equal devices and for CPU plus GPU, of the share of runs with 1,000 tries that find a near-best split, at 12, 16, 20 "
                       "and 29 qubits. The QAOA line falls from above 90 percent at 12 qubits to below 1 percent at 29 qubits. Local search and simulated annealing "
                       "also fall but stay higher; spectral bisection stays at 50 percent or more.",
    "fig_samples.png": f"Horizontal bar chart on a log scale of tries per optimum: random guessing {N['random_tries_per_optimum']:,.0f}, simulated annealing "
                       f"{N['annealing_median_evaluations']:,.0f}, local search {N['local_search_median_evaluations']:,.0f}, IonQ noise model {N['ionq_shots_per_optimum']:.1f}, "
                       f"noiseless {N['noiseless_shots_per_optimum']:.1f}.",
    "fig_qollab_placeholder.png": "Placeholder for a live Qollab screenshot. It shows lines of the recorded console output: running on ionq_simulator, "
                                  f"1,000 shots, measured P(optimum) {p_ionq:.4f}, and the two optimal splits measured {N['top_counts'][0]} and {N['top_counts'][1]} times.",
}

# ---------------------------------------------------------------- slides ----
slides = []


def slide(kicker, title, elements, notes, title_size=36):
    frame = [text(0.6, 0.45, 12.0, kicker, size=18, colour=ACCENT), text(0.6, 0.85, 12.2, title, size=title_size, bold=True)]
    slides.append(dict(elements=frame + elements, notes=notes))


slide("TEAM ZOLLERQU", "Splitting a simulation mesh\nbetween a CPU and a GPU", [
    text(0.6, 2.75, 7.4, "One qubit per mesh region. Run on IonQ's noisy\nsimulator and graded against the exact best split.", size=24, colour=MUTED),
    text(0.6, 4.25, 7.0, "Oussama Turki  ·  Tushar Pandey", size=24, bold=True),
    text(0.6, 4.78, 7.0, "Qollab x IonQ Hackathon, October 2026", size=18, colour=MUTED),
    text(0.6, 5.85, 7.4, ATTRIBUTION, size=20),
    card(8.45, 1.2, 4.3, 1.6), text(8.75, 1.38, 3.8, "12 qubits", size=36, colour=WHITE, bold=True),
    text(8.75, 2.08, 3.8, "one per mesh region", size=18),
    card(8.45, 3.0, 4.3, 1.6), text(8.75, 3.18, 3.8, f"{N['ionq_shots']:,} shots", size=36, colour=WHITE, bold=True),
    text(8.75, 3.88, 3.8, "IonQ Forte 1 noise model", size=18),
    card(8.45, 4.8, 4.3, 1.6), text(8.75, 4.98, 3.8, pct(p_ionq), size=36, colour=HIGHLIGHT, bold=True),
    text(8.75, 5.68, 3.8, "are the exact best split", size=18),
    source(8.45, 6.5, 4.4, f"IonQ Forte 1 noise model (Qollab), {date}"),
], notes="Hi, we are team ZollerQu: Oussama Turki and Tushar Pandey. We split a simulation mesh between two computers, a CPU and a GPU, with a quantum circuit on IonQ's noisy simulator.", title_size=40)

slide("THE PROBLEM", "One mesh, two devices: who computes what?", [
    text(0.6, 1.95, 3.0, pct(N["target_share"], 0), size=66, colour=WHITE, bold=True),
    text(3.25, 2.1, 3.9, "of the work should go\nto a GPU that is 4x\nfaster than the CPU", size=20),
    source(0.6, 3.35, 6.4, "assumed device speeds (model)"),
    text(0.6, 3.95, 6.6, "Who: engineers who run large simulations\n(heat, stress, crash) on CPU + GPU clusters.", size=20),
    text(0.6, 4.95, 6.6, "Why: a bad split leaves one device idle,\nor makes the devices exchange more data.", size=20),
    image("fig_problem.png", 7.5, 1.55, 5.3, 4.75, ALT["fig_problem.png"]),
    source(7.5, 6.33, 5.4, f"our mesh ({N['triangles']:,} triangles), the exact best split\nfor CPU + GPU (classical)"),
], notes='Big simulations, like crash tests, cut their mesh into parts. A good split gives each device a fair share of work and keeps the border short, because data crosses it at every step. A GPU four times faster should get eighty percent.')

slide("FORMULATION", "From mesh to Ising energy", [
    image("fig_formula.png", 0.9, 1.65, 11.5, 2.3, ALT["fig_formula.png"]),
    card(0.6, 4.3, 3.85, 2.05), text(0.85, 4.5, 3.4, "1 qubit =\n1 mesh region", size=24, colour=WHITE, bold=True),
    text(0.85, 5.55, 3.4, "0: device A   1: device B", size=18),
    card(4.74, 4.3, 3.85, 2.05), text(4.99, 4.5, 3.4, f"12 regions =\n{N['n_splits']:,} splits", size=24, colour=WHITE, bold=True),
    text(4.99, 5.55, 3.4, "all checked: the best is known", size=18),
    card(8.88, 4.3, 3.85, 2.05), text(9.13, 4.5, 3.4, "alpha =\ntarget work share", size=24, colour=WHITE, bold=True),
    text(9.13, 5.55, 3.4, "0.5 equal, 0.8 CPU + GPU", size=18),
    source(0.6, 6.45, 12.0, "exact best split: classical check of all splits"),
], notes='We group the mesh into twelve regions, one qubit each: zero means device A, one means device B. The energy adds the border that is cut and a penalty when the work share misses its target. Twelve qubits let us check every split and know the exact best one.')

slide("QUANTUM METHOD", "Fixed circuits, trained offline", [
    image("fig_circuit.png", 0.45, 1.7, 7.0, 4.3, ALT["fig_circuit.png"]),
    text(7.9, 1.85, 2.2, str(N["n_candidates"]), size=66, colour=WHITE, bold=True),
    text(10.0, 2.05, 2.9, "circuits compared\nat 12 qubits", size=20),
    text(7.9, 3.35, 5.0, "LR-QAOA, optimised QAOA, CVaR-QAOA,\nwarm starts, varQITE (IonQ's method)", size=18, colour=MUTED),
    text(7.9, 4.35, 5.0, "Trained on a noiseless simulator", size=22),
    text(7.9, 4.95, 5.0, "ONE fixed circuit per IonQ job", size=22, bold=True),
    source(7.9, 5.65, 5.0, "noiseless simulation (results/ionq_selection)"),
    source(0.45, 6.1, 7.0, "sketch of the QAOA circuit; varQITE uses the same start and measurement"),
], notes="The circuit starts from all splits, applies a cost layer with one two-qubit gate per pair of regions, then a mixer. We compared one hundred and four circuits, including IonQ's varQITE. Training is offline and noiseless. On Qollab we send one fixed circuit per job.")

slide("WHY IONQ", "Every region pairs with every other", [
    text(0.6, 1.95, 2.2, str(g12["all_to_all"]), size=66, colour=WHITE, bold=True),
    text(2.55, 2.1, 3.6, "two-qubit gates for one\nlayer at 12 qubits,\nall-to-all connectivity", size=20),
    text(0.6, 3.6, 5.6, f"{g12['heavy_hex']:.0f} after routing on a heavy-hex\nlayout: {g12['ratio']:.1f}x more", size=20),
    text(0.6, 4.6, 5.6, f"At 29 qubits: {g29['all_to_all']} against {g29['heavy_hex']:,.0f}\n({g29['ratio']:.1f}x more)", size=20),
    source(0.6, 5.65, 5.6, "counted with the Qiskit transpiler only\n(no simulation, no hardware run)"),
    image("fig_gates.png", 6.35, 1.6, 6.5, 4.5, ALT["fig_gates.png"]),
], notes='Our balance term couples every pair of regions. IonQ connects every pair of qubits, so one layer needs sixty-six two-qubit gates. Routed on a heavy-hex layout it needs about three times more, counted with the transpiler only.')

slide("MAIN RESULT  ·  2 EQUAL DEVICES", f"{100 * p_ionq:.0f} % of shots are the exact best split", [
    text(0.6, 1.8, 5.6, pct(p_ionq), size=80, colour=HIGHLIGHT, bold=True),
    text(0.6, 3.2, 5.8, f"of {N['ionq_shots']:,} shots, about {gain:,.0f}x random guessing", size=20),
    text(0.6, 3.68, 5.8, "Most frequent answer: the exact optimum", size=20),
    source(0.6, 4.2, 5.8, f"IonQ Forte 1 noise model (Qollab), {date}:\nvarQITE ZY, 12 qubits, {N['zy_two_qubit_gates']} two-qubit gates"),
    text(0.6, 4.95, 5.8, f"Noiseless: {pct(p_ideal)}. Noise costs {loss_points:.0f} points,\nabout {loss_points / estimated_loss:.0f}x what we estimated.", size=20),
    text(0.6, 5.8, 5.8, f"Classical, same 1,000 tries: annealing finds it\nin {pct(N['annealing_hit_1000'], 0)} of runs, local search in {pct(N['local_search_hit_1000'], 0)}.", size=18, colour=MUTED),
    image("fig_ionq_result.png", 6.55, 1.55, 6.3, 4.7, ALT["fig_ionq_result.png"]),
    source(6.55, 6.35, 6.3, "random and noiseless: noiseless simulation; classical: classical"),
], notes='The main result, for two equal devices. On the IonQ Forte 1 noise model, twenty-eight percent of one thousand shots were the exact best split, and it was the most frequent answer. Random guessing gives five hundredths of a percent. Without noise we reach seventy-five percent: noise costs twice what we estimated.')

q = N["q29"]
slide("SCALING  ·  HONEST COMPARISON", "Bigger meshes: classical methods win", [
    image("fig_scaling.png", 0.4, 1.55, 7.7, 4.6, ALT["fig_scaling.png"]),
    source(0.45, 6.2, 7.7, "server simulation, noiseless, 10 Oct 2026 (cases 2 and 9, seeds 7-9);\nQAOA = our best setting per size; the other lines are classical"),
    image("fig_samples.png", 8.25, 1.55, 4.7, 3.3, ALT["fig_samples.png"]),
    source(8.3, 4.9, 4.7, "12 qubits: IonQ Forte 1 noise model (Qollab),\nclassical"),
    text(8.3, 5.5, 4.7, f"Caveats: varQITE first needed {N['zy_training_circuits']:,}\nnoiseless training circuits. A shot is not\nthe same work as a classical evaluation.", size=18),
], notes="Now the honest part. As the mesh grows, classical methods win: at twenty-nine regions, simulated without noise, our best circuit finds a near-best split in under one percent of runs, spectral bisection in half or more. On IonQ's simulator we needed four shots per optimum against eleven hundred classical evaluations. But varQITE first needed seven thousand six hundred training circuits, and a shot is not the same work as a classical evaluation.")

slide("DEMO  ·  LIMITS  ·  NEXT STEPS", "Run it on Qollab", [
    image("fig_qollab_placeholder.png", 0.5, 1.6, 5.75, 3.5, ALT["fig_qollab_placeholder.png"]),
    text(0.6, 5.2, 5.8, "Qollab project: [pending]      Video: [pending]", size=18),
    text(0.6, 5.62, 5.8, "GitHub: Oussema-t/zollerqu-qollab (branch mesh)", size=18),
    text(6.7, 1.6, 6.2, "Limits", size=24, colour=ACCENT, bold=True),
    text(6.7, 2.1, 6.2, "1. A noisy simulator, not hardware.\n2. CPU + GPU: the varQITE circuit fails. One\n    qubit stays 50/50, and we proved why.\n"
                        "3. Beyond 12 qubits, classical methods win.", size=18),
    text(6.7, 3.95, 6.2, "Next steps", size=24, colour=ACCENT, bold=True),
    text(6.7, 4.45, 6.2, "1. Run the other circuits on the noise model.\n2. Repair the circuit for unequal shares.\n3. Split between more than two devices.", size=18),
    text(0.6, 6.3, 12.0, ATTRIBUTION, size=20),
], notes='[CUT TO THE LIVE QOLLAB DEMO] On Qollab we press Run and choose IonQ Forte 1. Limits: a noisy simulator, not hardware. For a CPU plus a GPU the varQITE circuit fails, and we proved why. Beyond twelve qubits classical methods win. Next, we repair that circuit. Thank you. This effort is supported by Qollab and IonQ.')

rows = [("Scenario", "Circuit (two-qubit gates)", "Random", "Noiseless", "IonQ noise model"),
        ("2 equal devices", f"varQITE ZY ({N['zy_two_qubit_gates']})", pct(p_random, 3), pct(p_ideal), pct(p_ionq)),
        ("CPU + GPU", f"CVaR-QAOA, p = 2 ({N['cpu_gpu_best_quantum_gates']})", pct(N["cpu_gpu_random_p_opt"], 3), pct(N["cpu_gpu_best_quantum_p_opt"], 2), "[pending]"),
        ("CPU + GPU", f"varQITE ZY ({N['zy_two_qubit_gates']})", pct(N["cpu_gpu_random_p_opt"], 3), "0.00 % (cap proved)", "[pending]")]
assert N["cpu_gpu_varqite_zy_p_opt"] < 5e-5 and N["cpu_gpu_best_quantum"].startswith("CVaR-QAOA")
columns = [(0.85, 2.2), (3.05, 3.6), (6.65, 1.3), (7.95, 2.55), (10.55, 2.15)]
table = []
for r, row in enumerate(rows):
    y = 1.9 + r * 0.95
    table.append(card(0.6, y - 0.18, 12.15, 0.8))
    for (x, w), cell in zip(columns, row):
        highlighted = r == 1 and cell == pct(p_ionq)
        table.append(text(x, y, w, cell, size=18, bold=r == 0 or highlighted, colour=ACCENT if r == 0 else HIGHLIGHT if highlighted else WHITE))
slide("APPENDIX", "Both device scenarios at 12 qubits", table + [
    text(0.6, 5.85, 12.0, "Share of shots that are the exact best split. The noisy runs for CPU + GPU are still to be made.", size=18, colour=MUTED),
    source(0.6, 6.35, 12.0, f"random, noiseless: noiseless simulation; last column: IonQ Forte 1 noise model (Qollab), {date}"),
], notes="Appendix, not spoken: the two device scenarios side by side. For CPU plus GPU the published varQITE circuit is capped by a qubit that stays "
         "fifty-fifty; the best other circuit reaches about half a percent without noise.")

# ---------------------------------------------------------------- checks: speaker notes length ----
spoken = " ".join(s["notes"] for s in slides[:8])
words = len(spoken.replace("[CUT TO THE LIVE QOLLAB DEMO]", "").split())
assert 330 <= words <= 380, f"speaker notes have {words} words (target 330-380)"


# ---------------------------------------------------------------- renderer 1: PDF with matplotlib (also measures overflow) ----
def render_pdf(path):
    overflow = []
    with PdfPages(path) as pdf:
        for number, s in enumerate(slides, start=1):
            fig = plt.figure(figsize=(W, H), facecolor=BG)
            ax = fig.add_axes([0, 0, 1, 1])
            ax.set_xlim(0, W)
            ax.set_ylim(H, 0)
            ax.axis("off")
            renderer = fig.canvas.get_renderer()
            footer = [text(0.6, 7.03, 8.0, FOOTER, size=12, colour=MUTED), text(11.7, 7.03, 1.03, str(number), size=12, colour=MUTED, align="r")]
            for e in s["elements"] + footer:
                if e["kind"] == "card":
                    ax.add_patch(FancyBboxPatch((e["x"], e["y"]), e["w"], e["h"], boxstyle="round,pad=0,rounding_size=0.12", facecolor=PANEL, edgecolor="none"))
                elif e["kind"] == "image":
                    picture = Image.open(e["path"]).convert("RGB")
                    scale = min(e["w"] / picture.width, e["h"] / picture.height)
                    w, h = picture.width * scale, picture.height * scale
                    x0, y0 = e["x"] + (e["w"] - w) / 2, e["y"] + (e["h"] - h) / 2
                    picture.thumbnail((1800, 1800))
                    ax.imshow(picture, extent=(x0, x0 + w, y0 + h, y0), aspect="auto", zorder=2)
                else:
                    for index, line in enumerate(e["lines"]):
                        x = e["x"] if e["align"] == "l" else e["x"] + e["w"]
                        artist = ax.text(x, e["y"] + index * e["size"] * LINE / 72, line, fontsize=e["size"], color=e["colour"], va="top",
                                         ha="left" if e["align"] == "l" else "right", fontfamily=e["font"], fontweight="bold" if e["bold"] else "normal",
                                         fontstyle="italic" if e["italic"] else "normal", zorder=3)
                        width = artist.get_window_extent(renderer).width / fig.dpi
                        if width > e["w"] + 0.02:
                            overflow.append(f"slide {number}: '{line[:50]}' is {width:.2f} in wide, box {e['w']:.2f} in")
                        if e["y"] + (index + 1) * e["size"] * LINE / 72 > H - 0.15 or e["x"] + min(width, e["w"]) > W - 0.3:
                            overflow.append(f"slide {number}: '{line[:50]}' leaves the slide")
            ax.set_xlim(0, W)
            ax.set_ylim(H, 0)
            pdf.savefig(fig, facecolor=BG)
            plt.close(fig)
    return overflow


# ---------------------------------------------------------------- renderer 2: PPTX with python-pptx ----
def rgb(colour):
    return RGBColor.from_string(colour.lstrip("#"))


def render_pptx(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(W), Inches(H)
    deck.core_properties.title = "ZollerQu - mesh partitioning on IonQ"
    deck.core_properties.author = "Oussama Turki, Tushar Pandey (Team ZollerQu)"
    for number, s in enumerate(slides, start=1):
        page = deck.slides.add_slide(deck.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = rgb(BG)
        footer = [text(0.6, 7.03, 8.0, FOOTER, size=12, colour=MUTED), text(11.7, 7.03, 1.03, str(number), size=12, colour=MUTED, align="r")]
        for e in s["elements"] + footer:
            if e["kind"] == "card":
                shape = page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                shape.adjustments[0] = min(0.5, 0.12 / min(e["w"], e["h"]))
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(PANEL)
                shape.line.fill.background()
                shape.shadow.inherit = False
            elif e["kind"] == "image":
                with Image.open(e["path"]) as picture:
                    scale = min(e["w"] / picture.width, e["h"] / picture.height)
                    w, h = picture.width * scale, picture.height * scale
                shape = page.shapes.add_picture(str(e["path"]), Inches(e["x"] + (e["w"] - w) / 2), Inches(e["y"] + (e["h"] - h) / 2), Inches(w), Inches(h))
                shape._element.xpath(".//p:cNvPr")[0].set("descr", e["alt"])
            else:
                height = len(e["lines"]) * e["size"] * LINE / 72 + 0.05
                box = page.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(height))
                frame = box.text_frame
                frame.word_wrap = False
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                frame.vertical_anchor = MSO_ANCHOR.TOP
                for index, line in enumerate(e["lines"]):
                    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    paragraph.alignment = PP_ALIGN.LEFT if e["align"] == "l" else PP_ALIGN.RIGHT
                    paragraph.line_spacing = Pt(e["size"] * LINE)
                    run = paragraph.add_run()
                    run.text = line
                    run.font.name, run.font.size, run.font.bold, run.font.italic = e["font"], Pt(e["size"]), e["bold"], e["italic"]
                    run.font.color.rgb = rgb(e["colour"])
        page.notes_slide.notes_text_frame.text = s["notes"]
    deck.save(path)


if __name__ == "__main__":
    problems = render_pdf(SLIDES / "mesh_ZollerQu.pdf")
    render_pptx(SLIDES / "mesh_ZollerQu.pptx")
    names = ["Title", "Problem", "Formulation", "Quantum method", "Why IonQ", "Main IonQ result", "Scaling and honest comparison",
             "Demo, limits and next steps", "Appendix: both device scenarios"]
    notes = ["# Speaker notes: mesh partitioning on IonQ (Team ZollerQu)", "",
             f"Spoken script for slides 1-8: {words} words, about 2:30-2:50 at a calm pace. The appendix is not spoken.", ""]
    for number, (name, s) in enumerate(zip(names, slides), start=1):
        notes += [f"## Slide {number}: {name}", "", s["notes"], ""]
    notes += ["## Figure alt text", ""] + [f"- **{name}**: {alt}  \n  Source: {SOURCES[name]}" for name, alt in ALT.items()]
    notes += ["", "## Placeholders still open", "", "- Slide 8: Qollab project link `[pending]`, video link `[pending]`, live Qollab screenshot `[pending]` "
              "(the picture shows the recorded console output).", "- Appendix: the IonQ Forte 1 noise-model runs for CPU + GPU `[pending]` (two cells)."]
    (SLIDES / "speaker_notes.md").write_text("\n".join(notes) + "\n")
    print(f"{len(slides)} slides written: mesh_ZollerQu.pptx, mesh_ZollerQu.pdf, speaker_notes.md ({words} spoken words)")
    print("text overflow: " + ("none" if not problems else ""))
    for problem in problems:
        print("  " + problem)
